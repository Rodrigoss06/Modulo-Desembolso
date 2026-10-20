"""Genera docs/Informe_Tecnico_Modulo_Desembolso.docx sobre la plantilla FPSI-teo05.docx.

Uso:  python docs/generar_informe.py
Inserta los diagramas de docs/imagenes/diagramas, las capturas de docs/imagenes/capturas
(si existen) y los resultados de docs/evidencias/pytest-junit.xml (si existe).
"""
from __future__ import annotations

import copy
import re
import struct
import xml.etree.ElementTree as ET
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
PLANTILLA = RAIZ / "docs/fuentes/FPSI-teo05.docx"
SALIDA = RAIZ / "docs/Informe_Tecnico_Modulo_Desembolso.docx"
DIAGRAMAS = RAIZ / "docs/imagenes/diagramas"
CAPTURAS = RAIZ / "docs/imagenes/capturas"
JUNIT = RAIZ / "docs/evidencias/pytest-junit.xml"
REPO = "https://github.com/Rodrigoss06/Modulo-Desembolso"

ANCHO_DXA = 9292  # ancho util de la seccion de contenido de la plantilla
ANCHO_CM = 16.39
ALTO_MAX_CM = 20.5

doc = Document(PLANTILLA)
cont = {"tabla": 0, "figura": 0}

# ---------------------------------------------------------------------------
# Utilidades de formato (replican el formato directo de la plantilla)
# ---------------------------------------------------------------------------
TOKENS = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))")


def _hipervinculo(p, url: str, texto: str, tam: float | None = None) -> None:
    r_id = doc.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(u)
    if tam:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(tam * 2)))
        rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = texto
    r.append(t)
    h.append(r)
    p._p.append(h)


def _runs(p, texto: str, tam: float | None = None, negrita: bool | None = None,
          cursiva: bool | None = None, color: str | None = None) -> None:
    """Agrega texto con marcado simple: **negrita**, `codigo`, [texto](url)."""
    for parte in TOKENS.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            r = p.add_run(parte[2:-2])
            r.bold = True
        elif parte.startswith("`") and parte.endswith("`"):
            r = p.add_run(parte[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt((tam or 12) - 1.5)
            if negrita:
                r.bold = True
            continue
        elif m := re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", parte):
            _hipervinculo(p, m.group(2), m.group(1), tam)
            continue
        else:
            r = p.add_run(parte)
            if negrita:
                r.bold = True
        if tam:
            r.font.size = Pt(tam)
        if cursiva:
            r.italic = True
        if color:
            r.font.color.rgb = RGBColor.from_string(color)


def _insertar_ppr(p, elemento) -> None:
    ppr = p._p.get_or_add_pPr()
    sucesores = {qn(x) for x in ("w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")}
    for i, hijo in enumerate(ppr):
        if hijo.tag in sucesores:
            ppr.insert(i, elemento)
            return
    ppr.append(elemento)


def parrafo(texto: str = "", *, alinear=WD_ALIGN_PARAGRAPH.JUSTIFY, antes: float = 0,
            despues: float = 6, interlineado: float = 1.0, tam: float | None = None,
            negrita: bool | None = None, cursiva: bool | None = None, color: str | None = None,
            mantener: bool = False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = alinear
    pf.space_before = Pt(antes)
    pf.space_after = Pt(despues)
    pf.line_spacing = interlineado
    if mantener:
        pf.keep_with_next = True
    if texto:
        _runs(p, texto, tam=tam, negrita=negrita, cursiva=cursiva, color=color)
    return p


def _nivel_esquema(p, nivel: int) -> None:
    o = OxmlElement("w:outlineLvl")
    o.set(qn("w:val"), str(nivel))
    _insertar_ppr(p, o)


def titulo1(texto: str) -> None:
    p = parrafo(texto, alinear=WD_ALIGN_PARAGRAPH.LEFT, antes=12, despues=7,
                interlineado=1.15, tam=14, negrita=True, mantener=True)
    _nivel_esquema(p, 0)


def titulo2(texto: str) -> None:
    p = parrafo(texto, alinear=WD_ALIGN_PARAGRAPH.LEFT, antes=14, despues=5,
                interlineado=1.15, tam=12, negrita=True, mantener=True)
    _nivel_esquema(p, 1)


def subtitulo(texto: str) -> None:
    parrafo(texto, despues=3, negrita=True, mantener=True)


def vinetas(items: list[str]) -> None:
    for item in items:
        p = parrafo(item, alinear=WD_ALIGN_PARAGRAPH.LEFT, despues=3)
        num = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), "2")
        num.append(ilvl)
        num.append(num_id)
        ppr = p._p.get_or_add_pPr()
        ppr.insert(0, num)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.5)


def salto_pagina() -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _celda(celda, texto: str, ancho: int, tam: float, encabezado: bool, sombra: str | None) -> None:
    tcpr = celda._tc.get_or_add_tcPr()
    tcw = OxmlElement("w:tcW")
    tcw.set(qn("w:w"), str(ancho))
    tcw.set(qn("w:type"), "dxa")
    tcpr.append(tcw)
    if sombra:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), sombra)
        tcpr.append(shd)
    mar = OxmlElement("w:tcMar")
    for lado, valor in (("top", 50), ("left", 80), ("bottom", 50), ("right", 80)):
        m = OxmlElement(f"w:{lado}")
        m.set(qn("w:w"), str(valor))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tcpr.append(mar)
    lineas = texto.split("\n") if texto else [""]
    primero = celda.paragraphs[0]
    for i, linea in enumerate(lineas):
        p = primero if i == 0 else celda.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(1)
        pf.line_spacing = 1.0
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _runs(p, linea, tam=tam, negrita=True if encabezado else None)


def tabla(titulo: str, encabezados: list[str], filas: list[list[str]], anchos_cm: list[float],
          tam: float = 10, nota: str | None = None) -> None:
    cont["tabla"] += 1
    p = parrafo(f"**Tabla {cont['tabla']}.** {titulo}", alinear=WD_ALIGN_PARAGRAPH.LEFT,
                antes=4, despues=3, tam=10, mantener=True)
    total = sum(anchos_cm)
    anchos = [round(a / total * ANCHO_DXA) for a in anchos_cm]
    anchos[-1] = ANCHO_DXA - sum(anchos[:-1])
    t = doc.add_table(rows=1 + len(filas), cols=len(encabezados))
    tbl = t._tbl
    tblpr = tbl.tblPr
    for hijo in list(tblpr):
        tblpr.remove(hijo)
    estilo = OxmlElement("w:tblStyle")
    estilo.set(qn("w:val"), "Tablaconcuadrcula")
    tblpr.append(estilo)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(ANCHO_DXA))
    tw.set(qn("w:type"), "dxa")
    tblpr.append(tw)
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "000000")
        bordes.append(b)
    tblpr.append(bordes)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    grid = tbl.tblGrid
    for col, ancho in zip(grid.findall(qn("w:gridCol")), anchos):
        col.set(qn("w:w"), str(ancho))
    for r, fila in enumerate([encabezados] + filas):
        tr = t.rows[r]._tr
        trpr = tr.get_or_add_trPr()
        if r == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        trpr.append(OxmlElement("w:cantSplit"))
        for c, texto in enumerate(fila):
            _celda(t.cell(r, c), texto, anchos[c], tam, r == 0, "D9D9D9" if r == 0 else None)
    if nota:
        parrafo(nota, alinear=WD_ALIGN_PARAGRAPH.LEFT, antes=2, despues=8, tam=9, cursiva=True)
    else:
        parrafo("", despues=4)


def codigo(texto: str) -> None:
    t = doc.add_table(rows=1, cols=1)
    tbl = t._tbl
    tblpr = tbl.tblPr
    for hijo in list(tblpr):
        tblpr.remove(hijo)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(ANCHO_DXA))
    tw.set(qn("w:type"), "dxa")
    tblpr.append(tw)
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "A6A6A6")
        bordes.append(b)
    tblpr.append(bordes)
    tbl.tblGrid.findall(qn("w:gridCol"))[0].set(qn("w:w"), str(ANCHO_DXA))
    celda = t.cell(0, 0)
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    tcpr.append(shd)
    for i, linea in enumerate(texto.strip("\n").split("\n")):
        p = celda.paragraphs[0] if i == 0 else celda.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(linea if linea else " ")
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
    parrafo("", despues=4)


def _dimensiones_png(ruta: Path) -> tuple[int, int]:
    with open(ruta, "rb") as f:
        cabecera = f.read(24)
    return struct.unpack(">II", cabecera[16:24])


def figura(ruta: Path, titulo: str, ancho_cm: float = ANCHO_CM) -> None:
    cont["figura"] += 1
    p = parrafo(alinear=WD_ALIGN_PARAGRAPH.CENTER, antes=4, despues=2, mantener=True)
    if ruta.exists():
        w, h = _dimensiones_png(ruta)
        ancho = min(ancho_cm, ALTO_MAX_CM * w / h)
        p.add_run().add_picture(str(ruta), width=Cm(ancho))
    else:
        _runs(p, f"[Captura pendiente: {ruta.name}]", tam=10, cursiva=True, color="7F7F7F")
    parrafo(f"**Figura {cont['figura']}.** {titulo}", alinear=WD_ALIGN_PARAGRAPH.CENTER,
            despues=10, tam=10)


def indice() -> None:
    p = doc.add_paragraph()
    r = p.add_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r._r.append(inicio)
    r._r.append(instr)
    r._r.append(sep)
    r2 = p.add_run("Clic derecho y «Actualizar campo» para generar el índice.")
    r2.italic = True
    r3 = p.add_run()
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    r3._r.append(fin)


def resultados_pruebas() -> dict[str, str]:
    if not JUNIT.exists():
        return {}
    estados = {}
    for caso in ET.parse(JUNIT).getroot().iter("testcase"):
        nombre = caso.get("name", "").split("[")[0]
        if caso.find("failure") is not None or caso.find("error") is not None:
            estados[nombre] = "Fallida"
        elif caso.find("skipped") is not None:
            estados.setdefault(nombre, "Omitida")
        else:
            estados.setdefault(nombre, "Aprobada")
    return estados


# ---------------------------------------------------------------------------
# Portada: se conserva la de la plantilla y se actualizan curso, tema y equipo
# ---------------------------------------------------------------------------
cuerpo = doc.element.body
hijos = list(cuerpo.iterchildren())
for el in hijos[9:]:
    if el.tag != qn("w:sectPr"):
        cuerpo.remove(el)

parrafos_portada = [p for p in doc.paragraphs if p.text.strip()]
p_curso, p_tema = parrafos_portada[0], parrafos_portada[1]
p_curso.runs[0].text = "GESTIÓN DE SISTEMAS Y TECNOLOGÍAS DE LA INFORMACIÓN - A"
for r in p_curso.runs[1:]:
    r.text = ""
p_tema.runs[0].text = "PROYECTO INTEGRADOR BNA – FASE II: "
p_tema.runs[1].text = "INFORME TÉCNICO DEL MÓDULO 13 DESEMBOLSO (LÍNEA DE CRÉDITOS)"
for r in p_tema.runs[2:]:
    r.text = ""

INTEGRANTES = [
    ("2023600951", "Centeno Lopez, Jose Alfredo"),
    ("2023223991", "Nuñez Arista, Piero Bruker"),
    ("2023224941", "Reaño Soto, Iván Edgardo"),
    ("2023225471", "Sucapuca Santamarina, Rodrigo Alonso"),
]
FECHA = date.today().strftime("%d/%m/%Y")
t_alumnos = doc.tables[1]
fila_modelo = t_alumnos.rows[1]._tr
for _ in INTEGRANTES[1:]:
    t_alumnos._tbl.append(copy.deepcopy(fila_modelo))
for fila, (codigo_alumno, nombre) in zip(t_alumnos.rows[1:], INTEGRANTES):
    for celda, valor in zip(fila.cells, (codigo_alumno, nombre, FECHA)):
        runs = celda.paragraphs[0].runs
        runs[0].text = valor
        for r in runs[1:]:
            r.text = ""

parrafo("")
parrafo("Índice", alinear=WD_ALIGN_PARAGRAPH.LEFT, antes=6, despues=6, tam=14, negrita=True)
indice()
salto_pagina()

# ---------------------------------------------------------------------------
# 1. Introducción
# ---------------------------------------------------------------------------
titulo1("1. Introducción")
parrafo("El módulo de Desembolso convierte un crédito aprobado por Evaluación en un abono "
        "confirmado en una cuenta de ahorros del titular, previa validación de su identidad, y "
        "comunica el resultado a Cobranza para que la obligación continúe su ciclo. Es el cuarto "
        "eslabón de la Línea de Créditos del Banco Nacional de Arequipa (BNA): Captación, "
        "Solicitud, Evaluación, Desembolso y Cobranza.")
parrafo("Este informe documenta el diseño técnico y la implementación del módulo en la Fase II: "
        "alcance, reglas de negocio, requerimientos funcionales y no funcionales, arquitectura, base "
        "de datos, interfaces, seguridad, pantallas, tecnologías, operación y pruebas. Parte del "
        "Informe TOGAF de la Fase I, del Modelo Referencial de Base de Datos, de la ficha del "
        "servicio SRV-CR-04 (I-01), de las prácticas ITIL 4 (I-04), del catálogo de niveles de "
        "servicio (I-05) y de los acuerdos AC-40 a AC-47 del acta ACT-2026-007.")
p = parrafo("El código fuente, los scripts de base de datos, los documentos fuente y este informe "
            "se encuentran en el repositorio del módulo: ")
_hipervinculo(p, REPO, REPO)
p.add_run(".")
parrafo("**Criterio aplicado.** Las decisiones que siguen pendientes ante otros comités "
        "(permisos, ejecutor financiero, calendario de pagos, política de reversión y contratos "
        "definitivos) no se presentan como aprobadas. El módulo las implementa como supuestos de "
        "trabajo explícitos y configurables, registrados en la sección 18, conforme a AC-40 y AC-47.")
tabla("Datos generales del módulo", ["Campo", "Detalle"], [
    ["Proyecto", "Plataforma Digital Banco Nacional de Arequipa (BNA)"],
    ["Línea y comité", "Línea de Créditos · Comité de Créditos (líder: Mathias Cabrera; enlace del módulo: Sebastián Conde)"],
    ["Módulo", "13. Desembolso (DES)"],
    ["Servicio de TI", "SRV-CR-04 · Servicio de Gestión del Desembolso de Créditos (código propuesto)"],
    ["Líder de equipo", "Rodrigo Alonso Sucapuca Santamarina (ACT-2026-007)"],
    ["Responsable funcional propuesto", "Iván Edgardo Reaño Soto (I-04 / I-05)"],
    ["Repositorio", f"[{REPO}]({REPO})"],
    ["Docentes", "Jorge Martínez · Mario Santillana"],
    ["Versión del informe", f"1.0 · {FECHA}"],
], [4.2, 12.2])

# ---------------------------------------------------------------------------
# 2. Alcance
# ---------------------------------------------------------------------------
titulo1("2. Alcance del módulo")
titulo2("2.1. Funciones incluidas")
vinetas([
    "Consultar los créditos con aprobación habilitada por Evaluación (CR-INT-04).",
    "Preparar la orden de desembolso tomando el importe y la moneda de la aprobación, nunca del canal.",
    "Validar la identidad del titular con el servicio de Clientes, sin almacenar biometría.",
    "Consultar y validar la cuenta de ahorros destino con Cuentas (CR-INT-05).",
    "Confirmar y ejecutar el desembolso con revalidación, idempotencia y control de concurrencia.",
    "Tratar los resultados inciertos (estado POR_VERIFICAR) consultando el resultado antes de cualquier otra acción.",
    "Entregar el desembolso confirmado a Cobranza mediante una bandeja de salida con reintentos (CR-INT-06).",
    "Consultar estado, historial, bitácora y constancia de cada orden.",
    "Conciliar por fecha las órdenes propias con los abonos de Cuentas y los acuses de Cobranza.",
    "Registrar solicitudes de reversión con autorización de un segundo supervisor, deshabilitadas por defecto hasta que se apruebe la política.",
])
titulo2("2.2. Escenario base")
parrafo("Desembolso total, gestionado en agencia, a una cuenta de ahorros individual habilitada del "
        "titular y en la misma moneda del crédito (PEN o USD). La elección acota el prototipo y es "
        "coherente con la dependencia Desembolso–Cuentas reconocida en la matriz transversal.")
titulo2("2.3. Fuera de alcance")
vinetas([
    "Desembolsos parciales, conversión de moneda, cuentas mancomunadas, comisiones e impuestos.",
    "Doble autorización del desembolso y desembolso por Home Banking (ese canal corresponde al pago de cuotas en Cobranza).",
    "Cálculo y persistencia del calendario de pagos (PC-01) y anulación de aprobaciones de Evaluación.",
])
titulo2("2.4. Responsabilidades externas")
tabla("Responsabilidades que el módulo no asume", ["Módulo", "Responsabilidad"], [
    ["12. Evaluación", "Decisión de aprobación, su versión y su vigencia."],
    ["11. Solicitud de Crédito", "Condiciones del crédito y calendario de pagos."],
    ["2. Clientes", "Validación de identidad y datos maestros del cliente."],
    ["3. Cuentas", "Titularidad, estado y elegibilidad de la cuenta; ejecución del abono y su asiento (CR-INT-05)."],
    ["14. Cobranza", "Activación de la obligación y cobro de cuotas."],
    ["9. Seguridad", "Autenticación OAuth2, emisión de tokens, antifraude y auditoría central."],
    ["8. Bus de Servicios", "Transporte, enrutamiento y versionado de los contratos REST/JSON."],
], [4.2, 12.2])

# ---------------------------------------------------------------------------
# 3. Posición y dependencias
# ---------------------------------------------------------------------------
titulo1("3. Posición en la Línea de Créditos y dependencias")
parrafo("Desembolso recibe una decisión ya tomada por Evaluación y entrega a Cobranza un crédito "
        "efectivamente desembolsado. No reabre decisiones anteriores ni administra el cobro posterior. "
        "Toda comunicación con otros módulos se modela como consumo de servicios a través del Bus, sin "
        "acceso a tablas ajenas (AC-40).")
figura(DIAGRAMAS / "01_contexto.png", "Posición del módulo en la Línea de Créditos y matriz CR-INT.", 11.5)
tabla("Interfaces que utiliza el módulo", ["Interfaz", "Sentido", "Uso en el módulo", "Estado del contrato"], [
    ["CR-INT-04", "Evaluación → Desembolso", "Consultar aprobaciones habilitadas y revalidarlas antes de ejecutar.", "Relación ratificada; contrato simulado"],
    ["CR-INT-05", "Desembolso → Cuentas", "Cuentas elegibles, elegibilidad, abono y consulta del abono por clave.", "Ejecutor financiero pendiente (OBS-DES-03)"],
    ["CR-INT-06", "Desembolso → Cobranza", "Eventos desembolso.confirmado y desembolso.revertido.", "Activación y calendario pendientes (OBS-DES-05)"],
    ["CR-INT-10", "Desembolso → Seguridad / Auditoría", "Registro de eventos críticos.", "Contrato transversal pendiente"],
    ["IF-02 (Fase I)", "Desembolso → Clientes", "Validación de identidad del titular.", "Por conciliar con la matriz (OBS-DES-08)"],
    ["IF-05 (Fase I)", "Seguridad → Desembolso", "Tokens OAuth2, JWKS y evaluación antifraude.", "Flujo y permisos por acordar (OBS-DES-02)"],
], [2.4, 3.6, 6.4, 4.0])

# ---------------------------------------------------------------------------
# 4. Actores, roles y permisos
# ---------------------------------------------------------------------------
titulo1("4. Actores, roles y permisos")
parrafo("Los códigos de permiso provienen del catálogo de control de accesos del BNA "
        "(`DESEMBOLSO_LEER` y `DESEMBOLSO_EJECUTAR`, asignados a `LIDER_DESEMBOLSO`). Los roles de "
        "operación en agencia son una propuesta del equipo, pendiente de ratificación con Seguridad "
        "(OBS-DES-02, REQ-01). El token emitido por Seguridad transporta el rol y los permisos.")
tabla("Roles del módulo", ["Rol", "Permisos", "Acciones permitidas"], [
    ["OPERADOR_AGENCIA", "DESEMBOLSO_LEER\nDESEMBOLSO_EJECUTAR", "Consultar créditos, preparar, validar, confirmar, consultar resultado y constancia, solicitar reversión."],
    ["SUPERVISOR", "DESEMBOLSO_LEER\nDESEMBOLSO_EJECUTAR", "Lo anterior, más monitor operativo, reintentos de entrega, conciliación, bitácora y decisión sobre reversiones solicitadas por otro usuario."],
    ["AUDITOR", "DESEMBOLSO_LEER", "Consultas, bitácora y conciliación, sin ejecutar operaciones."],
    ["Servicio del módulo", "Token de servicio (client credentials)", "Llamadas a otros módulos por el Bus y trabajos de recuperación y entrega."],
], [3.6, 4.4, 8.4])

# ---------------------------------------------------------------------------
# 5. Reglas de negocio
# ---------------------------------------------------------------------------
titulo1("5. Reglas de negocio")
parrafo("Las reglas RN-01 a RN-10 provienen del Informe TOGAF de la Fase I. La columna de origen "
        "distingue las reglas propias, los acuerdos de la Línea de Créditos y las reglas de Ahorros "
        "aplicables solo por interacción con Cuentas.")
tabla("Reglas de negocio y su implementación", ["ID", "Regla", "Origen / estado", "Implementación"], [
    ["RN-01", "La decisión de aprobación pertenece a Evaluación. Desembolso verifica su habilitación y no la modifica.", "Acuerdo de línea · Vigente", "Se consulta la aprobación al preparar y se revalida al confirmar; nunca se escribe en Evaluación."],
    ["RN-02", "Debe existir validación satisfactoria del cliente y autorización del operador antes de ejecutar.", "Propia · Vigente", "Validaciones IDENTIDAD y AUTORIZACION obligatorias y permiso DESEMBOLSO_EJECUTAR."],
    ["RN-03", "La cuenta destino cumple titularidad, estado y restricciones informadas por Cuentas.", "Ahorros por interacción · Por confirmar", "Elegibilidad consultada a Cuentas al seleccionar la cuenta y otra vez al ejecutar."],
    ["RN-04", "Importe y moneda corresponden al crédito aprobado; sin conversiones ni descuentos.", "Propia · Vigente", "El importe no se recibe del canal: se copia de la aprobación. La cuenta debe estar en la moneda del crédito."],
    ["RN-05", "Solo un desembolso total por crédito; una nueva clave técnica no elude el control.", "Propia · Vigente", "Índice único sobre `credito_en_ejecucion` y verificación previa al preparar."],
    ["RN-06", "La elegibilidad y los límites se revalidan en la frontera de ejecución.", "Propia · Vigente", "Validaciones de fase EJECUCION antes de pasar a EN_PROCESO."],
    ["RN-07", "Un tiempo de espera agotado no prueba que la operación haya fallado.", "Propia · Vigente", "Estado POR_VERIFICAR, consulta por la misma clave y ningún reintento del abono."],
    ["RN-08", "Una falla de entrega a Cobranza no convierte el desembolso en rechazado.", "Acuerdo de línea · Por confirmar", "`entrega_estado` independiente y outbox con reintentos."],
    ["RN-09", "Una reversión es otra operación financiera referenciada y auditable; no borra la original.", "Propia · Por confirmar", "`des_solicitud_reversion` con clave propia; la orden pasa a REVERTIDA conservando su historial."],
    ["RN-10", "El abono se realiza sobre una cuenta de ahorros administrada por Cuentas.", "Matriz transversal · Vigente", "Solo se aceptan cuentas de ahorros individuales que Cuentas declara elegibles."],
], [1.3, 5.4, 3.3, 6.4])
titulo2("5.1. Reglas técnicas derivadas")
tabla("Reglas técnicas propuestas por el equipo", ["ID", "Regla"], [
    ["RT-01", "Toda operación de negocio exige `X-Idempotency-Key`. Misma clave y mismo contenido devuelven el resultado existente; misma clave con contenido distinto devuelve conflicto 409."],
    ["RT-02", "Concurrencia optimista con la columna `version`: la confirmación exige la versión que vio el operador."],
    ["RT-03", "Importes como `Decimal` y `DECIMAL(15,2)`, nunca como flotantes; en JSON se transmiten como cadena."],
    ["RT-04", "Tiempo de espera de 3 s por llamada a dependencias; reintentos (máximo 2, con espera creciente) solo en consultas, nunca en el abono."],
    ["RT-05", "Las validaciones tienen vigencia; una validación vencida devuelve la orden a PREPARADA."],
    ["RT-06", "Una respuesta NO_ENCONTRADO de Cuentas solo es definitiva tras la ventana de gracia configurada."],
    ["RT-07", "Cuentas y documentos se enmascaran en pantallas, bitácora y registros (por ejemplo, ****4521)."],
    ["RT-08", "La reversión requiere un autorizador distinto del solicitante (control de cuatro ojos)."],
], [1.5, 14.9])

# ---------------------------------------------------------------------------
# 6. Requerimientos funcionales
# ---------------------------------------------------------------------------
titulo1("6. Requerimientos funcionales")
tabla("Requerimientos funcionales y su implementación", ["ID", "Requerimiento", "Prioridad", "Reglas", "Implementación"], [
    ["RF-01", "Obtener una aprobación habilitada para desembolso.", "Alta", "RN-01", "`GET /creditos-habilitados`, `POST /ordenes`"],
    ["RF-02", "Validar la identidad del cliente titular.", "Alta", "RN-02", "`POST /ordenes/{id}/identidad`"],
    ["RF-03", "Consultar y validar la cuenta destino.", "Alta", "RN-03, RN-10", "`GET /ordenes/{id}/cuentas-elegibles`, `PUT /ordenes/{id}/cuenta-destino`"],
    ["RF-04", "Preparar y confirmar la orden de desembolso.", "Alta", "RN-04, RN-06", "`POST /ordenes`; pantallas 2 y 3"],
    ["RF-05", "Ejecutar el desembolso mediante el Bus.", "Alta", "RN-04", "`POST /ordenes/{id}/confirmacion`"],
    ["RF-06", "Evitar duplicados y consultar operaciones inciertas.", "Alta", "RN-05, RN-07", "Idempotencia, índice único, `POST /ordenes/{id}/verificacion`, recuperador"],
    ["RF-07", "Comunicar el resultado del desembolso a Cobranza.", "Alta", "RN-08", "`des_evento_entrega` y despachador"],
    ["RF-08", "Consultar estado, historial y constancia.", "Media", "RN-07", "`GET /ordenes`, `GET /ordenes/{id}`, `GET /ordenes/{id}/constancia`"],
    ["RF-09", "Registrar y conciliar las operaciones ejecutadas.", "Media", "RN-05, RN-08", "`POST /conciliaciones`, vistas `des_v_*`"],
    ["RF-10", "Solicitar la reversión conforme a una política aprobada.", "Baja", "RN-09", "`POST /ordenes/{id}/reversiones`, `POST /reversiones/{id}/decision`"],
], [1.3, 4.8, 1.7, 2.2, 6.4], tam=9.5, nota="Todas las rutas cuelgan de /api/v1.")
tabla("Entradas, salidas y criterios de aceptación", ["ID", "Entradas", "Salidas", "Criterio de aceptación"], [
    ["RF-01", "Documento del cliente o identificador del crédito.", "Aprobación con versión, importe, moneda y vigencia.", "Una aprobación vencida, revocada o ya desembolsada impide preparar la orden."],
    ["RF-02", "Orden y método de verificación.", "Referencia de validación y resultado.", "La orden conserva una referencia de validación aprobada del titular; no se guarda biometría."],
    ["RF-03", "Orden, cuenta y moneda.", "Titularidad, estado, moneda y restricciones.", "Solo se continúa con una cuenta de ahorros individual, habilitada y en la moneda del crédito."],
    ["RF-04", "Aprobación, cuenta destino y usuario.", "Orden VALIDADA con clave de idempotencia.", "Se muestran crédito, cuenta enmascarada, moneda e importe y se registra al operador."],
    ["RF-05", "Orden validada, clave y correlación.", "Resultado con operación financiera y fecha efectiva.", "Una orden CONFIRMADA siempre tiene `operacion_financiera_id` y `fecha_efectiva`."],
    ["RF-06", "Clave de idempotencia y huella SHA-256 del contenido.", "Resultado existente, conflicto o estado por verificar.", "Repetir la misma solicitud devuelve el resultado existente; contenido distinto se rechaza."],
    ["RF-07", "Desembolso confirmado.", "Evento entregado y acuse de Cobranza.", "El evento se conserva hasta su acuse y puede reenviarse sin duplicar la activación."],
    ["RF-08", "Orden, crédito o cliente.", "Estado, referencias y constancia.", "Una persona autorizada conoce el resultado sin ejecutar nuevamente."],
    ["RF-09", "Fecha y moneda.", "Reporte de diferencias.", "Se comparan orden, abono en Cuentas y acuse de Cobranza; un origen incompleto no produce conciliación definitiva."],
    ["RF-10", "Orden confirmada, motivo y autorización.", "Reversión trazable y su resultado.", "La reversión referencia la operación original y conserva motivo, solicitante, autorizador y resultado."],
], [1.3, 4.2, 4.6, 6.3], tam=9.5)

# ---------------------------------------------------------------------------
# 7. Casos de uso
# ---------------------------------------------------------------------------
titulo1("7. Casos de uso")
tabla("Casos de uso del módulo", ["CU", "Actor", "Flujo principal", "Excepciones"], [
    ["CU-01 Consultar créditos habilitados", "Operador", "Busca por documento o crédito; el módulo consulta a Evaluación y muestra importe, moneda, estado y vigencia.", "E1 sin resultados. E2 aprobación vencida o revocada: se informa y no se continúa. E3 Evaluación no responde: 503 y registro del intento."],
    ["CU-02 Validar cliente y cuenta destino", "Operador, titular", "Valida identidad con Clientes, lista cuentas con Cuentas y selecciona la cuenta destino.", "E1 identidad rechazada: la orden sigue PREPARADA. E2 sin cuenta elegible. E3 moneda distinta: bloqueo del escenario base."],
    ["CU-03 Confirmar y ejecutar", "Operador", "Resumen, confirmación, revalidación, EN_PROCESO, abono en Cuentas, resultado y evento a Cobranza.", "E1 revalidación fallida: RECHAZADA sin invocar a Cuentas. E2 duplicidad: resultado existente. E3 sin respuesta: POR_VERIFICAR. E4 rechazo de Cuentas."],
    ["CU-04 Consultar resultado y constancia", "Operador", "Busca la orden; si está incierta consulta a Cuentas por la misma clave; la constancia exige confirmación.", "E1 orden inexistente: 404 sin datos de terceros. E2 sin permiso: 403 auditado."],
    ["CU-05 Entregar a Cobranza", "Servicio del módulo", "Toma eventos pendientes, los envía y registra el acuse.", "E1 Cobranza no disponible: reintento con espera creciente. E2 reenvío: Cobranza deduplica por evento. E3 se superan los intentos: FALLIDO y alerta."],
    ["CU-06 Conciliar operaciones", "Supervisor, auditor", "Compara órdenes, abonos de Cuentas y acuses de Cobranza de una fecha.", "E1 un origen no responde: no se emite conciliación definitiva."],
    ["CU-07 Solicitar reversión", "Operador, supervisor", "Solicitud con motivo; otro supervisor autoriza; se solicita a Cuentas con clave propia.", "E1 política no aprobada (deshabilitada). E2 autorizador igual al solicitante. E3 rechazo de Cuentas."],
], [3.3, 2.2, 5.6, 5.3], tam=9.5)

# ---------------------------------------------------------------------------
# 8. Requerimientos no funcionales
# ---------------------------------------------------------------------------
titulo1("8. Requerimientos no funcionales")
tabla("Requerimientos no funcionales", ["ID", "Requerimiento", "Meta o criterio", "Cómo se cumple"], [
    ["RNF-01", "Consistencia financiera", "Nunca existe una orden CONFIRMADA sin operación financiera; abono y asiento son responsabilidad de Cuentas/CORE.", "CHECK `ck_orden_confirmada`; CONFIRMADA solo con respuesta de Cuentas; evento outbox en la misma transacción."],
    ["RNF-02", "Seguridad y mínimo privilegio", "Toda operación autenticada y autorizada; los accesos denegados quedan auditados.", "JWT RS256 validado con JWKS; permiso por ruta; usuario de base de datos sin DELETE."],
    ["RNF-03", "Recuperación e idempotencia", "Ningún reintento produce un segundo abono.", "Clave de idempotencia, índice único por crédito, POR_VERIFICAR y recuperador."],
    ["RNF-04", "Auditoría", "Toda acción registrada con actor, resultado y `correlation_id`.", "`des_auditoria` de solo inserción protegida por triggers."],
    ["RNF-05", "Interoperabilidad", "Contratos REST/JSON versionados y documentados.", "Prefijo /api/v1; OpenAPI generado por FastAPI en /docs."],
    ["RNF-06", "Rendimiento", "p95 de la confirmación ≤ 3 s extremo a extremo (referencial, AC-42); tiempo propio medido aparte.", "Tiempo de espera de 3 s por dependencia; métricas de latencia por dependencia."],
    ["RNF-07", "Operabilidad", "Configuración por ambiente, salud, métricas y alertas.", "/health/live, /health/ready, /metrics y registros JSON con correlación."],
    ["RNF-08", "Mantenibilidad", "Reglas críticas cubiertas por pruebas automáticas.", "pytest con MySQL real y simulador de dependencias (sección 17)."],
    ["RNF-09", "Usabilidad", "Confirmar solo se habilita con validaciones completas; mensajes de rechazo accionables.", "Pantallas guiadas por el estado de la orden."],
], [1.5, 3.0, 5.6, 6.3], tam=9.5)

# ---------------------------------------------------------------------------
# 9. Arquitectura
# ---------------------------------------------------------------------------
titulo1("9. Arquitectura")
titulo2("9.1. Principios aplicados")
tabla("Principios de arquitectura (Fase I) y su aplicación", ["ID", "Principio", "Aplicación en el módulo"], [
    ["PA-01", "Integridad financiera", "El abono y su asiento los registra Cuentas/CORE; el módulo solo confirma con su respuesta."],
    ["PA-02", "Idempotencia", "Clave por operación de negocio; Cuentas y Cobranza deduplican por clave o evento."],
    ["PA-03", "Trazabilidad", "Orden, validaciones, operación financiera, evento y bitácora comparten `correlation_id`."],
    ["PA-04", "Propiedad de datos", "Ids externos como referencias lógicas sin FK; ninguna lectura de tablas ajenas."],
    ["PA-05", "Integración común", "Toda llamada sale por `BUS_BASE_URL` con REST/JSON y cabeceras estándar."],
    ["PA-06", "Decisiones justificadas", "Supuestos y pendientes registrados en la sección 18 con su configuración."],
    ["PA-07", "Seguridad por diseño", "Identidad, permiso y antifraude se comprueban antes de ejecutar."],
    ["PA-08", "Compatibilidad de línea", "Interfaces alineadas con la matriz CR-INT; nada se agrega sin coordinación."],
], [1.5, 3.6, 11.3])
titulo2("9.2. Vista de contenedores")
parrafo("La solución se compone de tres aplicaciones y la base de datos del módulo. El frontend en "
        "Next.js atiende al personal de agencia y sus Route Handlers actúan como BFF: guardan la sesión "
        "en una cookie httpOnly y llaman a la API, de modo que el navegador nunca conoce las URL "
        "internas. La API en FastAPI concentra las reglas, la persistencia y los trabajos en segundo "
        "plano. En desarrollo, un simulador del Bus de Servicios reproduce los contratos de Evaluación, "
        "Clientes, Cuentas, Cobranza y Seguridad, incluidas sus fallas.")
figura(DIAGRAMAS / "02_arquitectura.png", "Vista de contenedores del módulo de Desembolso.", 12.5)
titulo2("9.3. Componentes del backend")
figura(DIAGRAMAS / "03_componentes.png", "Componentes internos de la API (arquitectura por capas con adaptadores).", 9.0)
tabla("Responsabilidad de cada paquete", ["Paquete", "Responsabilidad"], [
    ["`app/api`", "Routers REST, esquemas Pydantic de entrada y salida, dependencias de seguridad."],
    ["`app/servicios`", "Casos de uso: preparar, validar, confirmar, verificar, consultar, entregar, recuperar, conciliar y revertir."],
    ["`app/dominio`", "Estados de la orden, transiciones válidas y errores de negocio."],
    ["`app/persistencia`", "Modelos SQLAlchemy de las tablas `des_*`, sesión y repositorio."],
    ["`app/integraciones`", "Clientes HTTP de Evaluación, Clientes, Cuentas, Cobranza y Seguridad a través del Bus."],
    ["`app/trabajos`", "Hilos del recuperador de órdenes inciertas y del despachador del outbox."],
    ["`app/core`", "Configuración, JWT y permisos, errores, correlación, UUID v7, dinero y enmascarado."],
], [4.0, 12.4])
titulo2("9.4. Proceso de desembolso")
figura(DIAGRAMAS / "04_proceso.png", "Proceso de desembolso en agencia y tratamiento de resultados.")
titulo2("9.5. Secuencia de confirmación")
parrafo("La confirmación guarda la intención (EN_PROCESO) y confirma la transacción antes de llamar a "
        "Cuentas. Así, una caída posterior siempre deja una orden recuperable y el índice único por "
        "crédito impide que otra orden llegue a Cuentas mientras tanto.")
figura(DIAGRAMAS / "05_secuencia_confirmacion.png", "Secuencia de confirmación y entrega a Cobranza.")
titulo2("9.6. Resultado incierto y recuperación")
vinetas([
    "Si la respuesta de Cuentas no llega (tiempo agotado, error 5xx o caída), la orden pasa a POR_VERIFICAR y el abono no se reenvía.",
    "El recuperador revisa cada `RECUPERACION_INTERVALO_SEG` las órdenes POR_VERIFICAR y las EN_PROCESO con más de `RECUPERACION_ESPERA_SEG`, y consulta `GET /cuentas/v1/abonos/{clave}`.",
    "CONFIRMADO lleva a CONFIRMADA y crea el evento; RECHAZADO lleva a RECHAZADA; NO_ENCONTRADO solo es definitivo pasada la ventana `RECUPERACION_GRACIA_SEG`; un error mantiene POR_VERIFICAR.",
    "Una orden incierta por más de `ALERTA_INCIERTA_MIN` minutos se reporta como alerta (incidente INC-01).",
    "El operador puede pedir la verificación inmediata desde la pantalla de resultado sin crear otra orden.",
])
titulo2("9.7. Entrega a Cobranza (outbox)")
vinetas([
    "El evento `desembolso.confirmado` se inserta en `des_evento_entrega` en la misma transacción que confirma la orden.",
    "El despachador toma eventos PENDIENTE con `SELECT … FOR UPDATE SKIP LOCKED`, de modo que varias instancias no envían el mismo evento a la vez.",
    "El envío usa el `id` del evento como clave de idempotencia; Cobranza responde con un acuse (`confirmacion_id`).",
    "Ante falla se reintenta con espera creciente (5 s × 2ⁿ, hasta 15 min); tras `ENTREGA_MAX_INTENTOS` el evento queda FALLIDO y visible en el monitor (incidente INC-02). El desembolso sigue CONFIRMADO.",
])

# ---------------------------------------------------------------------------
# 10. Base de datos
# ---------------------------------------------------------------------------
titulo1("10. Base de datos")
parrafo("La base del proyecto es la base unificada del BNA en MySQL 8.4 (InnoDB, utf8mb4). El "
        "módulo solo es dueño de las tablas `des_*`, definidas en `database/sql/08_desembolso.sql`; "
        "el script es autocontenido, de modo que el módulo puede desplegarse con su propio esquema. "
        "Los identificadores son UUID v7 en CHAR(36) generados por la aplicación.")
titulo2("10.1. Modelo de datos")
figura(DIAGRAMAS / "07_er_desembolso.png", "Modelo de datos del módulo (referencias a otros módulos sin llave foránea).", 11.0)
tabla("Tablas del módulo", ["Tabla", "Responsabilidad", "Claves y restricciones"], [
    ["`des_orden`", "Intención, ejecución y resultado del desembolso.", "UK código, UK clave de idempotencia, UK `credito_en_ejecucion`; CHECK de importe, moneda, cuenta, confirmación y rechazo."],
    ["`des_validacion`", "Evidencia de aprobación, identidad, cuenta, autorización y antifraude.", "FK a la orden; tipo y fase (PREPARACION o EJECUCION)."],
    ["`des_evento_entrega`", "Outbox hacia Cobranza.", "UK (orden, tipo); índice por estado y siguiente intento."],
    ["`des_auditoria`", "Bitácora de solo inserción.", "Triggers que bloquean UPDATE y DELETE; índice por correlación."],
    ["`des_solicitud_reversion`", "Reversiones y su autorización.", "UK de clave; una reversión activa por orden; autorizador distinto del solicitante."],
    ["`des_correlativo`", "Correlativo anual del código DES-AAAA-NNNNNN.", "PK por año; se incrementa con SELECT … FOR UPDATE."],
], [4.0, 5.4, 7.0])
titulo2("10.2. Estados de la orden")
parrafo("El trigger `trg_des_orden_bu` impone las transiciones de la figura y vuelve inmutables "
        "importe, moneda, cuenta, crédito, aprobación y clave desde que la orden inicia su ejecución. "
        "El estado de entrega a Cobranza es independiente (PENDIENTE o CONFIRMADA).")
figura(DIAGRAMAS / "06_estados_orden.png", "Transiciones válidas del estado de la orden.", 9.5)
titulo2("10.3. Integridad impuesta por la base")
tabla("Controles de integridad verificados en MySQL 8.4", ["Control", "Mecanismo", "Efecto"], [
    ["Un desembolso por crédito", "Columna generada `credito_en_ejecucion` + UNIQUE", "Una segunda orden del mismo crédito no puede pasar a EN_PROCESO (error 1062)."],
    ["Idempotencia", "UNIQUE `idempotency_key` + `request_hash`", "Una clave repetida nunca crea otra orden."],
    ["Máquina de estados", "Trigger BEFORE UPDATE", "Transiciones inválidas rechazadas (SQLSTATE 45000)."],
    ["Inmutabilidad financiera", "Trigger BEFORE UPDATE", "No se modifica importe, cuenta ni operación tras iniciar la ejecución."],
    ["Bitácora inalterable", "Triggers BEFORE UPDATE / DELETE", "`des_auditoria` solo admite inserciones."],
    ["Cuatro ojos", "CHECK `ck_reversion_autorizador`", "El autorizador de una reversión no puede ser el solicitante."],
    ["Mínimo privilegio", "`usuario_app_desembolso.sql`", "El usuario de la aplicación no tiene DELETE en ninguna tabla."],
], [3.6, 5.2, 7.6])
titulo2("10.4. Cambios respecto del Modelo Referencial v1.0")
tabla("Ajustes al modelo entregado en la Fase I", ["Cambio", "Motivo"], [
    ["`credito_confirmado` se reemplaza por `credito_en_ejecucion`, que cubre EN_PROCESO, POR_VERIFICAR, CONFIRMADA y REVERTIDA.", "Con solo CONFIRMADA, dos órdenes del mismo crédito podían llegar a Cuentas a la vez (RN-05)."],
    ["`core_operacion_id` pasa a `operacion_financiera_id`.", "En la Fase II el abono se solicita a Cuentas por CR-INT-05; el ejecutor final sigue pendiente."],
    ["`cuenta_id` admite nulo mientras la orden está PREPARADA.", "La cuenta se elige y valida después de crear la orden (CU-02)."],
    ["`des_validacion` agrega el tipo APROBACION y la columna `fase`.", "Evidencia de la revalidación en la frontera de ejecución (RN-06)."],
    ["La reversión agrega estados EN_PROCESO y POR_VERIFICAR, origen y autorizador.", "El resultado incierto también aplica al extorno (RN-07) y AC-44 distingue su origen."],
    ["Nueva tabla `des_correlativo` y triggers de integridad.", "Código legible sin carreras y defensa en profundidad."],
    ["Sin llaves foráneas hacia tablas de otros módulos.", "PA-04 y AC-40."],
], [8.4, 8.0])

# ---------------------------------------------------------------------------
# 11. Interfaces
# ---------------------------------------------------------------------------
titulo1("11. Interfaces")
titulo2("11.1. API que ofrece el módulo")
tabla("Rutas de la API de Desembolso", ["Método y ruta", "Permiso o rol", "Uso"], [
    ["GET /creditos-habilitados", "DESEMBOLSO_LEER", "CU-01: aprobaciones por documento o crédito."],
    ["POST /ordenes", "DESEMBOLSO_EJECUTAR + clave", "Prepara la orden (PREPARADA)."],
    ["GET /ordenes · GET /ordenes/{id}", "DESEMBOLSO_LEER", "Historial con filtros y detalle con validaciones."],
    ["POST /ordenes/{id}/identidad", "DESEMBOLSO_EJECUTAR", "Validación de identidad con Clientes."],
    ["GET /ordenes/{id}/cuentas-elegibles", "DESEMBOLSO_EJECUTAR", "Cuentas del titular informadas por Cuentas."],
    ["PUT /ordenes/{id}/cuenta-destino", "DESEMBOLSO_EJECUTAR", "Selecciona y valida la cuenta destino."],
    ["POST /ordenes/{id}/confirmacion", "DESEMBOLSO_EJECUTAR + clave", "Revalida y ejecuta el desembolso."],
    ["POST /ordenes/{id}/verificacion", "DESEMBOLSO_LEER", "Consulta inmediata de una orden incierta."],
    ["GET /ordenes/{id}/constancia", "DESEMBOLSO_LEER", "Constancia; solo con confirmación financiera."],
    ["GET /ordenes/{id}/auditoria", "SUPERVISOR o AUDITOR", "Bitácora de la orden."],
    ["GET /operacion/resumen", "SUPERVISOR o AUDITOR", "Órdenes inciertas y entregas pendientes."],
    ["POST /operacion/recuperacion · /operacion/entregas", "SUPERVISOR", "Ejecuta a demanda el recuperador o el despachador."],
    ["POST /operacion/entregas/{evento}/reintento", "SUPERVISOR", "Reprograma un evento FALLIDO."],
    ["POST /conciliaciones", "SUPERVISOR o AUDITOR", "Conciliación por fecha y moneda."],
    ["POST /ordenes/{id}/reversiones · POST /reversiones/{id}/decision · GET /reversiones", "EJECUTAR · SUPERVISOR distinto · LEER", "Reversión (solo si `REVERSION_HABILITADA=true`)."],
    ["GET /health/live · /health/ready · /metrics", "Sin token (red interna)", "Salud y métricas para operación."],
], [6.2, 4.2, 6.0], tam=9.5, nota="Las rutas de negocio cuelgan de /api/v1; OpenAPI disponible en /docs.")
titulo2("11.2. Cabeceras y errores")
parrafo("Toda petición lleva `Authorization: Bearer <JWT>` y `X-Correlation-ID` (si falta, la API la "
        "genera y la devuelve). Las operaciones que afectan saldo exigen `X-Idempotency-Key`. Los "
        "errores usan un formato único:")
codigo('{ "codigo": "CUENTA_NO_ELEGIBLE",\n  "mensaje": "La cuenta destino no está habilitada para recibir el desembolso.",\n  "detalle": { "motivo": "CUENTA_BLOQUEADA" },\n  "correlation_id": "0199f2a0-7c1e-7b51-9a3d-2f6c1b9e8a10" }')
tabla("Códigos de error principales", ["HTTP", "Código", "Situación"], [
    ["401 / 403", "NO_AUTENTICADO · PERMISO_DENEGADO", "Token ausente o inválido · permiso insuficiente (se audita)."],
    ["404", "ORDEN_NO_ENCONTRADA · APROBACION_NO_ENCONTRADA", "Recurso inexistente, sin revelar datos de terceros."],
    ["409", "IDEMPOTENCIA_CONFLICTO · VERSION_CONFLICTO · ESTADO_INVALIDO · CREDITO_CON_DESEMBOLSO", "Clave con otro contenido · versión desactualizada · acción no válida en el estado · crédito ya en ejecución o desembolsado."],
    ["422", "APROBACION_NO_HABILITADA · IDENTIDAD_RECHAZADA · CUENTA_NO_ELEGIBLE · MONEDA_DISTINTA · VALIDACIONES_INCOMPLETAS · VALIDACION_VENCIDA · REVALIDACION_FALLIDA", "Reglas de negocio no satisfechas."],
    ["409", "REVERSION_DESHABILITADA · AUTORIZADOR_INVALIDO", "Política no aprobada · cuatro ojos."],
    ["503", "DEPENDENCIA_NO_DISPONIBLE · CONCILIACION_INCOMPLETA", "Un módulo contraparte no respondió."],
], [1.7, 7.0, 7.7], tam=9)
titulo2("11.3. Contratos consumidos")
parrafo("Los contratos siguientes son la propuesta del equipo para cerrar CR-INT-04, CR-INT-05 y "
        "CR-INT-06. Están implementados en el simulador del Bus y deben ratificarse con cada módulo "
        "(OBS-DES-08); cualquier cambio se tramita como solicitud de cambio (AC-45, CHG-01).")
tabla("Operaciones consumidas a través del Bus", ["Módulo", "Operación", "Respuesta esperada"], [
    ["Evaluación", "GET /evaluacion/v1/aprobaciones?documento=&credito_id=\nGET /evaluacion/v1/aprobaciones/{id}", "Aprobación: id, versión, crédito, cliente, importe, moneda, estado (HABILITADA, VENCIDA, REVOCADA, DESEMBOLSADA) y vigencia."],
    ["Clientes", "POST /clientes/v1/validaciones-identidad", "Referencia opaca, resultado (APROBADA o RECHAZADA) y vencimiento."],
    ["Cuentas", "GET /cuentas/v1/cuentas-elegibles?cliente_id=&moneda=\nGET /cuentas/v1/cuentas/{id}/elegibilidad", "Cuentas enmascaradas con tipo, moneda, estado, titularidad y elegibilidad con motivo."],
    ["Cuentas", "POST /cuentas/v1/abonos (X-Idempotency-Key)\nGET /cuentas/v1/abonos/{clave}\nGET /cuentas/v1/abonos?fecha=", "CONFIRMADO con `operacion_id` y `fecha_efectiva`, RECHAZADO con motivo, o NO_ENCONTRADO."],
    ["Cuentas", "POST /cuentas/v1/reversiones · GET /cuentas/v1/reversiones/{clave}", "Reversión confirmada o rechazada, con clave propia."],
    ["Cobranza", "POST /cobranza/v1/eventos-desembolso (X-Idempotency-Key = evento)\nGET /cobranza/v1/eventos-desembolso?fecha=", "Acuse `confirmacion_id`; un reenvío devuelve el mismo acuse sin activar dos veces."],
    ["Seguridad", "POST /seguridad/v1/token · GET /seguridad/v1/jwks.json\nPOST /seguridad/v1/antifraude/evaluaciones · POST /seguridad/v1/eventos", "Tokens RS256, claves públicas, decisión antifraude y recepción de eventos críticos."],
], [2.2, 7.4, 6.8], tam=9)
titulo2("11.4. Evento desembolso.confirmado (CR-INT-06)")
codigo('{\n  "evento_id": "0199f2a1-1b0c-7e4d-8f21-5c3a9d0e7b42",\n  "tipo": "desembolso.confirmado",\n  "version_esquema": "v1",\n  "orden_id": "0199f2a0-9d3e-7a10-b5c2-1e8f4a6d2c90",\n  "codigo_orden": "DES-2026-000001",\n  "credito_id": "CRE-2026-0001",\n  "aprobacion_id": "APR-2026-0001",\n  "aprobacion_version": 1,\n  "cliente_id": "CLI-0001",\n  "cuenta_id": "CTA-0001",\n  "operacion_financiera_id": "OPF-2026-000001",\n  "importe": "10000.00",\n  "moneda": "PEN",\n  "fecha_efectiva": "2026-10-19T10:15:42"\n}')

# ---------------------------------------------------------------------------
# 12. Seguridad
# ---------------------------------------------------------------------------
titulo1("12. Seguridad")
vinetas([
    "**Autenticación:** el frontend obtiene un JWT del servicio de Seguridad (OAuth2) y lo guarda en una cookie httpOnly, SameSite=Lax; el navegador nunca accede al token.",
    "**Validación del token:** la API verifica firma RS256 con las claves públicas del JWKS de Seguridad, además de emisor, audiencia y expiración.",
    "**Autorización:** cada ruta exige un permiso (`DESEMBOLSO_LEER` o `DESEMBOLSO_EJECUTAR`) o un rol; un acceso denegado responde 403 y se registra en la bitácora.",
    "**Datos sensibles:** no se almacena biometría ni datos maestros de otros módulos; cuentas y documentos se muestran enmascarados.",
    "**Secretos:** se leen de variables de entorno; el repositorio solo contiene `.env.example`.",
    "**Base de datos:** usuario de aplicación de mínimo privilegio y bitácora inalterable.",
    "**Transporte:** HTTPS/TLS en Testing y Producción a través del Bus; CORS restringido al origen del frontend.",
])

# ---------------------------------------------------------------------------
# 13. Pantallas
# ---------------------------------------------------------------------------
titulo1("13. Pantallas")
parrafo("Las pantallas siguen los bocetos de la Fase I (pantallas 1 a 4) y agregan las vistas de "
        "consulta y operación. La interfaz se guía por el estado de la orden: confirmar solo se "
        "habilita con las validaciones completas, y mientras la orden está en proceso o por "
        "verificar solo se permite consultar, nunca crear otra orden.")
tabla("Pantallas del módulo", ["Pantalla", "Ruta", "Campos y acciones"], [
    ["Inicio de sesión", "/login", "Usuario y contraseña contra Seguridad; muestra rol y agencia."],
    ["1. Consulta de créditos", "/creditos", "Búsqueda por documento o crédito; crédito, cliente, importe, moneda, estado y vigencia; preparar desembolso."],
    ["2. Preparación", "/ordenes/{id}", "Cliente y crédito con documento enmascarado; validación de identidad; selección de cuenta elegible; importe y moneda derivados."],
    ["3. Confirmación", "/ordenes/{id}", "Resumen con cuenta enmascarada, importe y moneda; volver o confirmar el desembolso."],
    ["4. Resultado", "/ordenes/{id}", "Estado financiero, operación y fecha efectiva, entrega a Cobranza; verificar resultado; constancia."],
    ["Constancia", "/ordenes/{id}/constancia", "Comprobante imprimible, solo con confirmación financiera."],
    ["Historial", "/ordenes", "Filtros por texto, estado y fechas; acceso al detalle."],
    ["Monitor operativo", "/operacion", "Órdenes inciertas y entregas pendientes o fallidas; ejecutar recuperación, entregas y reintentos."],
    ["Conciliación", "/conciliacion", "Fecha y moneda; totales y diferencias por origen."],
    ["Reversiones", "/reversiones", "Solicitudes y decisión de un supervisor distinto; aviso cuando la política está deshabilitada."],
], [3.4, 3.4, 9.6], tam=9.5)
parrafo("Las capturas siguientes se tomaron del sistema en ejecución (frontend, API, MySQL 8.4 y simulador "
        "del Bus) con Playwright, recorriendo los flujos reales: desembolso completo, respuesta perdida de "
        "Cuentas, caída de Cobranza, rechazo de Cuentas e identidad rechazada. Los datos son sintéticos.")
CAPTURAS_INFORME = [
    ("01_login.png", "Inicio de sesión contra el servicio de Seguridad."),
    ("02_consulta_creditos.png", "Pantalla 1. Consulta de créditos habilitados."),
    ("03_preparacion.png", "Pantalla 2. Preparación: identidad validada y selección de la cuenta destino."),
    ("04_confirmacion.png", "Pantalla 3. Confirmación con resumen y cuenta enmascarada."),
    ("05_resultado.png", "Pantalla 4. Resultado confirmado con operación financiera y entrega a Cobranza."),
    ("06_constancia.png", "Constancia del desembolso."),
    ("07_por_verificar.png", "Resultado incierto: orden POR_VERIFICAR con opción de verificar."),
    ("08_moneda_distinta.png", "Bloqueo de una cuenta en moneda distinta a la del crédito."),
    ("09_historial.png", "Historial de órdenes con filtros."),
    ("10_detalle_auditoria.png", "Detalle de la orden con validaciones y bitácora."),
    ("11_monitor.png", "Monitor operativo del supervisor."),
    ("12_conciliacion.png", "Conciliación diaria contra Cuentas y Cobranza."),
    ("13_swagger.png", "Documentación OpenAPI de la API (Swagger UI)."),
]
for archivo, titulo in CAPTURAS_INFORME:
    figura(CAPTURAS / archivo, titulo, 15.5)

# ---------------------------------------------------------------------------
# 14. Tecnologías
# ---------------------------------------------------------------------------
titulo1("14. Tecnologías")
tabla("Stack tecnológico", ["Capa", "Tecnología", "Uso"], [
    ["Backend", "Python 3.13 · FastAPI · Uvicorn", "API REST, OpenAPI y trabajos en segundo plano."],
    ["Persistencia", "MySQL 8.4 (InnoDB, utf8mb4) · SQLAlchemy 2.0 · PyMySQL", "Tablas `des_*`, transacciones ACID y bloqueos de fila."],
    ["Validación y configuración", "Pydantic 2 · pydantic-settings", "Contratos de entrada y salida; configuración por ambiente."],
    ["Integración", "httpx", "Llamadas REST al Bus con tiempos de espera y correlación."],
    ["Seguridad", "PyJWT · cryptography", "Validación de JWT RS256 con JWKS."],
    ["Frontend", "Next.js 16 (App Router) · React 19 · TypeScript 5", "Pantallas de agencia y Route Handlers como BFF."],
    ["Estilos y formularios", "Tailwind CSS 4 · React Hook Form · Zod", "Diseño responsivo y validación tipada en el cliente."],
    ["Simulación", "FastAPI (simulador del Bus)", "Contratos de Evaluación, Clientes, Cuentas, Cobranza y Seguridad con fallas configurables."],
    ["Pruebas", "pytest · Playwright", "Pruebas de reglas e integración; capturas del informe."],
    ["Documentación", "OpenAPI · Mermaid · python-docx", "Contratos, diagramas y este informe."],
], [3.4, 6.2, 6.8])

# ---------------------------------------------------------------------------
# 15. Repositorio, ambientes y ejecución
# ---------------------------------------------------------------------------
titulo1("15. Repositorio, ambientes y ejecución")
p = parrafo("Repositorio: ")
_hipervinculo(p, REPO, REPO)
tabla("Estructura del repositorio", ["Carpeta", "Contenido"], [
    ["`backend/`", "API FastAPI del módulo, pruebas y `.env.example`."],
    ["`simuladores/`", "Simulador del Bus de Servicios y de los módulos contraparte."],
    ["`frontend/`", "Aplicación Next.js para el personal de agencia."],
    ["`database/`", "Base BNA unificada con el script 08 del módulo, guía y scripts de carga."],
    ["`docs/`", "Este informe, su generador, diagramas, capturas y documentos fuente."],
], [3.6, 12.8])
tabla("Puertos en desarrollo", ["Servicio", "URL"], [
    ["Frontend", "http://localhost:3000"],
    ["API de Desembolso · OpenAPI", "http://localhost:8000 · http://localhost:8000/docs"],
    ["Simulador del Bus", "http://localhost:9000"],
    ["MySQL", "localhost:3306 (base `bna`)"],
], [6.4, 10.0])
parrafo("Pasos para levantar el entorno local (Windows, PowerShell):", despues=3)
codigo("""cd database; .\\cargar_bna.ps1            # crea la base bna (o solo sql\\08_desembolso.sql)
cd ..\\simuladores; python -m venv .venv; .venv\\Scripts\\pip install -r requirements.txt
.venv\\Scripts\\uvicorn app.main:app --port 9000
cd ..\\backend; python -m venv .venv; .venv\\Scripts\\pip install -r requirements-dev.txt
copy .env.example .env; .venv\\Scripts\\uvicorn app.main:app --port 8000
cd ..\\frontend; npm install; copy .env.example .env.local; npm run dev""")
tabla("Usuarios del simulador de Seguridad (solo desarrollo; contraseña Demo2026!)", ["Usuario", "Rol", "Uso"], [
    ["operador1 · operador2", "OPERADOR_AGENCIA", "Preparar, validar, confirmar y consultar."],
    ["supervisor1 · supervisor2", "SUPERVISOR", "Monitor, conciliación, reintentos y decisión de reversiones (cuatro ojos)."],
    ["auditor1", "AUDITOR", "Consultas, bitácora y conciliación sin ejecutar."],
], [4.4, 4.0, 8.0])
tabla("Ambientes", ["Ambiente", "Características"], [
    ["Desarrollo", "Datos sintéticos, simulador del Bus y MySQL local; reversión deshabilitada por defecto."],
    ["Testing", "Integración con los módulos reales a través del Bus; evidencias de I-06."],
    ["Producción", "Versiones aprobadas; secretos y usuario de base de datos administrados por Infraestructura."],
], [3.2, 13.2])

# ---------------------------------------------------------------------------
# 16. Operación y niveles de servicio
# ---------------------------------------------------------------------------
titulo1("16. Operación y niveles de servicio (ITIL 4)")
parrafo("El módulo expone los datos necesarios para medir los indicadores propuestos en I-05. Las "
        "metas numéricas siguen pendientes de validación con el Comité e Infraestructura (AC-42).")
tabla("Indicadores de I-05 y su fuente en el módulo", ["Indicador", "Cálculo", "Fuente"], [
    ["Disponibilidad", "(minutos acordados − minutos sin servicio) / minutos acordados × 100", "/health/ready y registros"],
    ["Tiempo de respuesta", "p95 de la confirmación, separando la espera en dependencias", "`des_latencia_dependencia_segundos` en /metrics"],
    ["Tasa de fallas", "órdenes RECHAZADAS por error técnico / intentos válidos × 100", "`des_ordenes_por_estado` y bitácora"],
    ["Abonos sin confirmar", "órdenes EN_PROCESO o POR_VERIFICAR / solicitudes de abono × 100", "`des_ordenes_inciertas` y vista `des_v_ordenes_inciertas`"],
    ["Entregas pendientes", "eventos no confirmados y su antigüedad", "`des_entregas_pendientes` y vista `des_v_entregas_pendientes`"],
], [3.4, 7.2, 5.8], tam=9.5)
tabla("Incidentes de I-04 y respuesta del módulo", ["Incidente", "Respuesta implementada"], [
    ["INC-01 Abono sin resultado confirmado de Cuentas", "POR_VERIFICAR, recuperador, verificación manual, alerta por antigüedad y prohibición de reintentar el abono."],
    ["INC-02 Falla de notificación a Cobranza", "Outbox con reintentos, estado FALLIDO visible en el monitor y reintento manual; el desembolso sigue CONFIRMADO."],
    ["Continuidad: caída de la API o MySQL", "Las órdenes quedan en estados recuperables; al volver, el recuperador y el despachador retoman sin duplicar."],
], [5.4, 11.0])

# ---------------------------------------------------------------------------
# 17. Pruebas
# ---------------------------------------------------------------------------
titulo1("17. Pruebas y verificación")
parrafo("Las pruebas automatizadas se ejecutan con pytest contra MySQL 8.4 real (script "
        "`08_desembolso.sql`) y contra el simulador del Bus levantado en un hilo, de modo que los "
        "tiempos de espera, la concurrencia y los triggers se ejercitan de verdad. Los escenarios "
        "replican la tabla de verificación de la Fase I.")
ESCENARIOS = [
    ("Identidad rechazada", "La orden no avanza y no se invoca a Cuentas.", "test_identidad_rechazada_no_avanza"),
    ("Cuenta no elegible", "No se habilita la confirmación.", "test_cuenta_no_elegible_no_habilita_confirmacion"),
    ("Cuenta en moneda distinta", "Se bloquea el escenario base.", "test_cuenta_moneda_distinta_bloqueada"),
    ("Aprobación vencida o revocada", "No se puede preparar la orden.", "test_aprobacion_no_habilitada_no_permite_preparar"),
    ("Preparación repetida con la misma clave", "Devuelve la misma orden.", "test_preparacion_idempotente"),
    ("Misma clave con contenido distinto", "Conflicto 409; la orden original no cambia.", "test_misma_clave_contenido_distinto_conflicto"),
    ("Flujo completo", "Orden CONFIRMADA con operación, evento y constancia.", "test_flujo_completo_confirmado"),
    ("Dos confirmaciones concurrentes del mismo crédito", "Una única operación financiera.", "test_confirmaciones_concurrentes_un_solo_abono"),
    ("Revalidación fallida al confirmar", "RECHAZADA sin invocar el abono.", "test_revalidacion_fallida_rechaza_sin_abono"),
    ("Cuentas rechaza el abono", "RECHAZADA con motivo; sin operación.", "test_rechazo_de_cuentas"),
    ("Cuentas confirma y se pierde la respuesta", "POR_VERIFICAR y luego CONFIRMADA sin duplicar el abono.", "test_respuesta_perdida_se_recupera_sin_duplicar"),
    ("Cobranza no disponible", "El desembolso sigue confirmado y la entrega pendiente.", "test_cobranza_caida_mantiene_confirmado"),
    ("Entrega fallida tras los reintentos", "Evento FALLIDO visible en el monitor; el reintento manual lo entrega.", "test_entrega_fallida_tras_reintentos_y_reintento_manual"),
    ("Reenvío del mismo evento", "Cobranza no activa dos veces el crédito.", "test_reenvio_evento_no_duplica_activacion"),
    ("Acceso sin permisos", "403 y registro en la bitácora.", "test_acceso_sin_permiso_es_auditado"),
    ("Constancia de orden no confirmada", "No se emite.", "test_constancia_solo_confirmada"),
    ("Conciliación del día", "Sin diferencias; origen caído produce 503.", "test_conciliacion"),
    ("Reversión", "Deshabilitada por defecto; con política, cuatro ojos.", "test_reversion"),
    ("Integridad en base de datos", "Triggers y restricciones rechazan cambios inválidos.", "test_bd_integridad"),
]
estados = resultados_pruebas()
filas = []
for escenario, esperado, prueba in ESCENARIOS:
    filas.append([escenario, esperado, f"`{prueba}`", estados.get(prueba, "Pendiente de ejecución")])
tabla("Escenarios de verificación y resultado", ["Escenario", "Resultado esperado", "Prueba", "Resultado"], filas,
      [4.0, 5.0, 5.4, 2.0], tam=9)
if estados:
    total = len(estados)
    aprobadas = sum(1 for v in estados.values() if v == "Aprobada")
    parrafo(f"Resultado de la última ejecución: {aprobadas} de {total} pruebas aprobadas "
            "(`docs/evidencias/pytest-junit.xml`).")

# ---------------------------------------------------------------------------
# 18. Supuestos y pendientes
# ---------------------------------------------------------------------------
titulo1("18. Supuestos de trabajo y decisiones pendientes")
tabla("Pendientes de la Fase II y tratamiento en la implementación", ["ID", "Tema pendiente", "Supuesto implementado", "Configuración"], [
    ["OBS-DES-02", "Permisos para consumir Cuentas y Evaluación", "El módulo llama con su token de servicio; el usuario solo necesita DESEMBOLSO_*.", "`SERVICIO_CLIENTE_ID`"],
    ["OBS-DES-03", "Ejecutor del abono y del asiento", "Cuentas ejecuta y devuelve `operacion_id`; si es el CORE, solo cambia el adaptador.", "`BUS_BASE_URL`"],
    ["OBS-DES-04", "Idempotencia y unicidad", "Clave por orden y unicidad por crédito; consulta antes de reintentar.", "`RECUPERACION_*`"],
    ["OBS-DES-05 / PC-01", "Activación y calendario", "Desembolso envía fecha efectiva y versión; no calcula el calendario.", "—"],
    ["OBS-DES-06", "Entrega durable a Cobranza", "Outbox con llamada idempotente y reintentos.", "`ENTREGA_*`"],
    ["OBS-DES-07 / PC-05", "Política de reversión", "Implementada con cuatro ojos y plazo, deshabilitada por defecto.", "`REVERSION_HABILITADA`, `REVERSION_PLAZO_HORAS`"],
    ["OBS-DES-08", "Contratos, errores y tiempos de espera", "Contratos de la sección 11.3 y tiempo de espera de 3 s.", "`TIMEOUT_*`"],
    ["OBS-DES-09", "Ampliaciones", "Fuera del escenario base; la moneda distinta se bloquea.", "—"],
    ["AC-42", "SLA definitivo", "Se mide solo el tiempo propio; los tiempos extremo a extremo son referenciales.", "/metrics"],
], [2.5, 3.8, 6.8, 3.3], tam=9)

# ---------------------------------------------------------------------------
# 19. Conclusiones
# ---------------------------------------------------------------------------
titulo1("19. Conclusiones")
vinetas([
    "El módulo implementa el servicio SRV-CR-04 dentro de los límites acordados: decide sobre su propia orden y consume como servicios la aprobación, la identidad, la cuenta, el abono y la entrega a Cobranza.",
    "La integridad financiera no depende de una sola capa: la API controla estados e idempotencia, y la base de datos repite los controles críticos (un desembolso por crédito, transiciones válidas, bitácora inalterable).",
    "El estado POR_VERIFICAR y el outbox convierten las fallas de comunicación en situaciones recuperables y trazables, en lugar de abonos duplicados o desembolsos perdidos para Cobranza.",
    "Las decisiones pendientes quedan aisladas en adaptadores y parámetros de configuración, de modo que su ratificación por los comités no obliga a rediseñar el módulo.",
])

# ---------------------------------------------------------------------------
# Referencias
# ---------------------------------------------------------------------------
titulo1("Referencias")
REFERENCIAS = [
    "AXELOS. (2019). ITIL Foundation: ITIL 4 edition. The Stationery Office.",
    "Comité de Créditos BNA. (2026). Referencia de dependencias TOGAF y matriz CR-INT. Universidad Católica de Santa María.",
    "Comité de Créditos BNA. (2026). Lineamiento técnico común de la Línea de Créditos: Python, FastAPI y MySQL. Universidad Católica de Santa María.",
    "Comité de Créditos BNA. (2026). Informe consolidado de Fase II ITIL 4 de la Línea de Créditos (v0.1). Universidad Católica de Santa María.",
    "Equipo del módulo de Desembolso. (2026). Informe de arquitectura del módulo 13 Desembolso, Fase I (TOGAF 10). Universidad Católica de Santa María.",
    "Equipo del módulo de Desembolso. (2026). Modelo referencial de base de datos del módulo 13 Desembolso (v1.0). Universidad Católica de Santa María.",
    "Equipo del módulo de Desembolso. (2026). Servicios de TI y valor generado del módulo Desembolso (I-01), prácticas ITIL 4 (I-04) y catálogo de niveles de servicio (I-05). Universidad Católica de Santa María.",
    "Equipo del módulo de Desembolso. (2026). Acta de reunión ACT-2026-007. Universidad Católica de Santa María.",
    "The Open Group. (2022). The TOGAF Standard, 10th Edition. The Open Group.",
]
for ref in REFERENCIAS:
    p = parrafo(ref, alinear=WD_ALIGN_PARAGRAPH.LEFT, despues=6)
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-1.0)

doc.core_properties.title = "Informe técnico del módulo 13 Desembolso"
doc.core_properties.author = "Rodrigo Alonso Sucapuca Santamarina"
doc.save(SALIDA)
print(f"{SALIDA.name}: {cont['tabla']} tablas, {cont['figura']} figuras")
